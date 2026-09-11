"""
Accuracy harness for the inspection report -> approval note flow.

Scores ConfigIQ against the hand-made ground truth for the two demo pump
reports (P-101, P-202), each in three forms: a PDF with a text layer, a
scanned PDF with no text layer, and a scanned PNG.

Stages
  extract  Read each report (text layer or OCR), pull out its values and check
           them against the SOP limits. Deterministic, no LLM.
  note     Run the whole agent with the local Ollama models and score the Word
           approval note it writes: the document an engineer would sign.

Usage
  python scripts/eval_inspection_accuracy.py --stage extract
  python scripts/eval_inspection_accuracy.py --stage note --out results/accuracy/note.json

The note stage only reads the finished .docx, so it scores any version of the
pipeline the same way.

Metrics
  fields       equipment tag, each reading's value and limit status, severity,
               and the inspector's recommendations
  silent       readings missing from the output with no warning to the engineer
  untraceable  figures in the note found neither in the report nor in the SOP
  assurance    FALSE when the note says "SOURCED & VERIFIED" while a value,
               status or severity is wrong or a figure is untraceable
"""
import argparse
import asyncio
import json
import os
import re
import sys
import time

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, ROOT)
os.chdir(ROOT)

INSPECTION_DIR = os.path.join("demo_data", "inspection")
SOP_PATH = os.path.join("demo_data", "knowledge_base", "SOP-M-104_centrifugal_pump_maintenance.md")
TASK = "Review this inspection report and prepare the approval note."

REPORTS = {"P-101": "P101", "P-202": "P202"}
FORMS = ["clean.pdf", "scanned.pdf", "scanned.png"]

GT_KIND = {
    "Vibration RMS": "vibration",
    "Bearing Temperature": "bearing_temperature",
    "Mechanical Seal Leakage": "seal_leakage",
    "Suction Pressure": "suction_pressure",
    "Discharge Pressure": "discharge_pressure",
    "Flow Rate": "flow_rate",
}
# Ground-truth status -> statuses ConfigIQ may report for it
GT_STATUS = {
    "EXCEEDED": {"EXCEEDED", "CRITICAL"},
    "WARNING": {"WARNING"},
    "NORMAL": {"OK"},
    "SLIGHT DROP": {"DEVIATION"},
}
# Words that express each status in free text, for scoring a Word note
STATUS_WORDS = {
    "EXCEEDED": ("exceeded", "critical", "over limit", "zone d", "high:"),
    "WARNING": ("warning", "zone c"),
    "NORMAL": ("normal", "within limit"),
    "SLIGHT DROP": ("below", "deviation", "slightly low"),
}
KIND_WORDS = {
    "vibration": "vibration",
    "bearing_temperature": "temperature",
    "seal_leakage": "leak",
    "suction_pressure": "suction",
    "discharge_pressure": "discharge",
    "flow_rate": "flow",
}
STOPWORDS = {"with", "from", "that", "this", "before", "after", "into", "each", "until", "within", "levels"}


def normalise(text: str) -> str:
    """Lower-case text and spell units one way, so values compare across sources."""
    t = (text or "").lower()
    t = re.sub(r"deg\s*c\b|degc\b|°\s*c\b|º\s*c\b", "°c", t)
    t = re.sub(r"drops\s*/\s*minute", "drops/min", t)
    t = re.sub(r"m3\s*/\s*h|m³\s*/\s*h", "m³/h", t)
    return t


VALUE_RE = re.compile(r"(?<![\w.])(\d+(?:\.\d+)?)\s*(mm/s|°c|drops/min|mpa|m³/h|kw|rpm|mm(?!/))")


def values_in(text: str) -> set:
    return {(round(float(n), 3), u) for n, u in VALUE_RE.findall(normalise(text))}


def norm_unit(unit: str) -> str:
    u = normalise(unit).strip()
    return {"drops/minute": "drops/min"}.get(u, u)


def pdf_text_layer(path: str) -> str:
    """A PDF's text layer, if it has one; scans and images have none."""
    if not path.lower().endswith(".pdf"):
        return ""
    import pypdf

    try:
        return "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(path).pages)
    except Exception:
        return ""


def load_gt(stem: str) -> dict:
    with open(os.path.join(INSPECTION_DIR, f"ground_truth_{stem}.json"), encoding="utf-8") as f:
        return json.load(f)


def rec_matches(gt_rec: str, lines: list) -> bool:
    """A ground-truth recommendation counts as present when most of its key words appear in one line."""
    words = [w for w in re.findall(r"[a-z]{4,}", gt_rec.lower()) if w not in STOPWORDS]
    if not words:
        return False
    for line in lines:
        low = line.lower()
        hits = sum(1 for w in words if w in low)
        if hits / len(words) >= 0.5:
            return True
    return False


# ---------------------------------------------------------------- extract stage

def score_assessment(a: dict, gt: dict) -> dict:
    fields = []

    def add(name, ok, got, want):
        fields.append({"field": name, "ok": bool(ok), "got": got, "want": want})

    tag = (a.get("equipment_tag") or {}).get("value")
    add("equipment_tag", tag == gt["equipment"], tag, gt["equipment"])
    date = (a.get("inspection_date") or {}).get("value")
    not_in_document = []
    # Some forms of a report omit the date; only score it where the page actually shows it.
    if gt["inspection_date"] in re.sub(r"\s+", "", a.get("text", "")):
        add("inspection_date", date == gt["inspection_date"], date, gt["inspection_date"])
    else:
        not_in_document.append("inspection_date")

    by_kind = {}
    for m in a.get("measurements", []):
        by_kind.setdefault(m["kind"], []).append(m)

    review_text = " ".join(a.get("review_items", [])).lower()
    silent = []
    for gm in gt["measurements"]:
        kind = GT_KIND[gm["parameter"]]
        cands = by_kind.get(kind, [])
        # Vibration is read at two bearings; the ground truth records the governing (highest) one.
        m = max(cands, key=lambda x: x["value"]) if cands else None
        got_val = f"{m['value']} {m['unit']}" if m else None
        value_ok = m is not None and abs(m["value"] - float(gm["value"])) < 1e-6 and norm_unit(m["unit"]) == norm_unit(gm["unit"])
        add(f"{kind}.value", value_ok, got_val, f"{gm['value']} {gm['unit']}")
        status_ok = m is not None and m.get("status") in GT_STATUS[gm["status"]]
        add(f"{kind}.status", status_ok, m.get("status") if m else None, gm["status"])
        if m is None and KIND_WORDS[kind] not in review_text:
            silent.append(kind)

    sev = (a.get("severity") or {}).get("value")
    add("severity", sev == gt["severity"], sev, gt["severity"])
    sops = a.get("sop_refs", [])
    add("governing_sops", set(gt["governing_sops"]) <= set(sops), sops, gt["governing_sops"])

    rec_lines = [r["text"] for r in a.get("recommendations", [])]
    for i, rec in enumerate(gt["recommendations"]):
        add(f"recommendation[{i}]", rec_matches(rec, rec_lines), None, rec)

    return {
        "fields": fields,
        "correct": sum(f["ok"] for f in fields),
        "total": len(fields),
        "silent_misses": silent,
        "not_in_document": not_in_document,
        "review_items": a.get("review_items", []),
        "text_source": a.get("text_source"),
    }


def run_extract() -> dict:
    from backend.documents.inspection_checks import assess_document

    cases = []
    for tag, stem in REPORTS.items():
        gt = load_gt(stem)
        for form in FORMS:
            path = os.path.join(INSPECTION_DIR, f"inspection_report_{stem}_{form}")
            t0 = time.time()
            assessment = assess_document(path)
            result = score_assessment(assessment, gt)
            result.update({"case": f"{tag} {form}", "seconds": round(time.time() - t0, 1)})
            cases.append(result)
    return {"stage": "extract", "cases": cases}


# ---------------------------------------------------------------- note stage

def read_note(path: str) -> dict:
    from docx import Document

    doc = Document(path)
    paras = [(p.style.name, p.text) for p in doc.paragraphs]
    rows = [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
    return {"paras": paras, "rows": rows, "lines": [t for _, t in paras] + rows}


def section_lines(paras: list, heading_word: str) -> list:
    out, inside = [], False
    for style, text in paras:
        if style.startswith("Heading"):
            inside = heading_word.lower() in text.lower()
            continue
        if inside and text.strip():
            out.append(text)
    return out


def score_note(docx_path: str, gt: dict, source_values: set) -> dict:
    note = read_note(docx_path)
    lines = note["lines"]
    full = normalise("\n".join(lines))
    fields = []

    def add(name, ok, got, want):
        fields.append({"field": name, "ok": bool(ok), "got": got, "want": want})

    add("equipment_tag", gt["equipment"].lower() in full, None, gt["equipment"])

    m = re.search(r"risk level:\s*([a-z]+)", full)
    sev = m.group(1).upper() if m else None
    add("severity", sev == gt["severity"], sev, gt["severity"])

    missing, silent = [], []
    review_lines = [normalise(l) for l in lines if "review" in l.lower()]
    for gm in gt["measurements"]:
        kind = GT_KIND[gm["parameter"]]
        want = (round(float(gm["value"]), 3), norm_unit(gm["unit"]))
        holding = [normalise(l) for l in lines if want in values_in(l)]
        add(f"{kind}.value", bool(holding), None, f"{gm['value']} {gm['unit']}")
        status_ok = any(any(w in l for w in STATUS_WORDS[gm["status"]]) for l in holding)
        add(f"{kind}.status", status_ok, None, gm["status"])
        if not holding:
            missing.append(kind)
            if not any(KIND_WORDS[kind] in l for l in review_lines):
                silent.append(kind)

    recs = section_lines(note["paras"], "Recommended Actions")
    for i, rec in enumerate(gt["recommendations"]):
        add(f"recommendation[{i}]", rec_matches(rec, recs), None, rec)

    untraceable = sorted(f"{v:g} {u}" for v, u in values_in("\n".join(lines)) if (v, u) not in source_values)
    claims_verified = "sourced & verified" in full
    # Recommendations are quoted from the report, so a ground-truth recommendation
    # worded differently does not make the note's assurance false.
    wrong = [f["field"] for f in fields if not f["ok"] and not f["field"].startswith("recommendation")]
    return {
        "fields": fields,
        "correct": sum(f["ok"] for f in fields),
        "total": len(fields),
        "silent_misses": silent,
        "untraceable_values": untraceable,
        "claims_verified": claims_verified,
        "false_assurance": claims_verified and bool(wrong or untraceable),
        "review_flagged": "human review required" in full,
    }


async def run_note() -> dict:
    os.environ.setdefault("LLM_PROVIDER", "ollama")
    from backend.agents.agent import ConfigIQAgent

    agent = ConfigIQAgent()
    with open(SOP_PATH, encoding="utf-8") as f:
        sop_text = f.read()
    cases = []
    for tag, stem in REPORTS.items():
        gt = load_gt(stem)
        with open(os.path.join(INSPECTION_DIR, f"inspection_report_{stem}_clean.txt"), encoding="utf-8") as f:
            report_text = f.read()
        for form in FORMS:
            path = os.path.abspath(os.path.join(INSPECTION_DIR, f"inspection_report_{stem}_{form}"))
            # The table-layout PDFs carry limits the narrative text lacks, so trace against the input too.
            source_values = values_in("\n".join([report_text, sop_text, pdf_text_layer(path)]))
            t0 = time.time()
            state = await agent.run(task=TASK, document_ids=[path], parameters={"file_path": path})
            docx = next((p for p in state.generated_files if p.endswith(".docx")), None)
            if docx:
                result = score_note(docx, gt, source_values)
            else:
                result = {"fields": [], "correct": 0, "total": 1, "error": state.error or "no approval note generated"}
            result.update({
                "case": f"{tag} {form}",
                "seconds": round(time.time() - t0, 1),
                "docx": os.path.basename(docx) if docx else None,
                "is_verified": state.is_verified,
            })
            cases.append(result)
            print(f"  done {tag} {form} in {result['seconds']}s", flush=True)
    return {"stage": "note", "cases": cases}


# ---------------------------------------------------------------- report

def summarise(report: dict) -> dict:
    cases = report["cases"]
    groups = {}
    for c in cases:
        for f in c["fields"]:
            key = f["field"].split(".")[-1].split("[")[0]
            g = groups.setdefault(key, [0, 0])
            g[0] += f["ok"]
            g[1] += 1
    summary = {
        "fields_correct": sum(c["correct"] for c in cases),
        "fields_total": sum(c["total"] for c in cases),
        "by_field": {k: f"{v[0]}/{v[1]}" for k, v in groups.items()},
        "silent_misses": sum(len(c.get("silent_misses", [])) for c in cases),
    }
    if report["stage"] == "note":
        summary["untraceable_values"] = sum(len(c.get("untraceable_values", [])) for c in cases)
        summary["false_assurance_notes"] = sum(bool(c.get("false_assurance")) for c in cases)
    summary["accuracy_pct"] = round(100.0 * summary["fields_correct"] / max(1, summary["fields_total"]), 1)
    return summary


def print_report(report: dict) -> None:
    print(f"\n{'case':<20} {'fields':>8} {'silent':>7}", end="")
    extra = report["stage"] == "note"
    print(f" {'untraceable':>12} {'assurance':>10}" if extra else "")
    for c in report["cases"]:
        print(f"{c['case']:<20} {c['correct']:>3}/{c['total']:<4} {len(c.get('silent_misses', [])):>7}", end="")
        if extra:
            assurance = "FALSE" if c.get("false_assurance") else ("VERIFIED" if c.get("claims_verified") else "REVIEW")
            print(f" {len(c.get('untraceable_values', [])):>12} {assurance:>10}")
        else:
            print()
    print("\nsummary:", json.dumps(report["summary"], indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--stage", choices=["extract", "note"], default="extract")
    parser.add_argument("--out", help="write the full JSON report here")
    args = parser.parse_args()

    report = run_extract() if args.stage == "extract" else asyncio.run(run_note())
    report["summary"] = summarise(report)
    print_report(report)
    if args.out:
        os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump(report, f, indent=2, default=str)
        print(f"\nsaved {args.out}")


if __name__ == "__main__":
    main()
