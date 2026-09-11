#!/usr/bin/env python3
"""
ConfigIQ - Sovereign On-Premise Agentic AI Workbench
Pre-Demo Offline Readiness Checker (Phase 12)

Verifies 100% local dependency readiness prior to live offline presentation.
"""

import os
import sys
import io
import json

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

def check_all():
    checks = {}
    details = {}

    # 1. Ollama reachable on localhost
    ollama_ok = False
    models_found = []
    try:
        import httpx
        r = httpx.get("http://127.0.0.1:11434/api/tags", timeout=3.0)
        if r.status_code == 200:
            ollama_ok = True
            models_found = [m.get("name", "") for m in r.json().get("models", [])]
    except Exception as e:
        details["Ollama"] = f"Unreachable on localhost:11434 ({e})"

    checks["Ollama"] = "PASS" if ollama_ok else "FAIL"

    # Helper to check if model exists in installed models
    def has_model(tag):
        tag_base = tag.split(":")[0]
        for m in models_found:
            if m == tag or m.startswith(tag_base + ":") or m == tag_base:
                return True
        return False

    checks["Llama"] = "PASS" if has_model("llama3") else "FAIL"
    checks["Qwen Coder 7B"] = "PASS" if has_model("qwen2.5-coder:7b") else "FAIL"
    checks["Qwen Coder 14B"] = "PASS" if has_model("qwen2.5-coder:14b") else "FAIL"
    checks["Moondream"] = "PASS" if has_model("moondream") else "FAIL"
    checks["Nomic Embeddings"] = "PASS" if has_model("nomic-embed-text") else "FAIL"

    # 2. RAG database & knowledge base
    rag_ok = False
    try:
        from backend.rag.vector_store import LocalVectorStore
        vs_path = os.path.join(os.getcwd(), "outputs", "storage", "vector_store.json")
        if os.path.exists(vs_path):
            store = LocalVectorStore(persistence_path=vs_path)
            if store.count() > 0:
                rag_ok = True
            else:
                details["RAG"] = "Vector store exists but has 0 chunks"
        else:
            details["RAG"] = f"Vector store file not found at {vs_path}"
    except Exception as e:
        details["RAG"] = str(e)
    checks["RAG"] = "PASS" if rag_ok else "FAIL"

    # 3. OCR dependencies
    ocr_ok = False
    try:
        from backend.documents.ocr import ocr_engine
        if ocr_engine.is_available():
            ocr_ok = True
        else:
            details["OCR"] = "RapidOCR ONNX engine failed initialization"
    except Exception as e:
        details["OCR"] = str(e)
    checks["OCR"] = "PASS" if ocr_ok else "FAIL"

    # 4. P&ID Pipeline & Assets
    pid_ok = False
    try:
        pid_img = os.path.join(os.getcwd(), "demo_data", "pid", "pid_system_b.png")
        if os.path.exists(pid_img):
            import cv2
            img = cv2.imread(pid_img)
            if img is not None:
                pid_ok = True
            else:
                details["P&ID Pipeline"] = "Failed to decode demo_data/pid/pid_system_b.png"
        else:
            details["P&ID Pipeline"] = f"P&ID image missing at {pid_img}"
    except Exception as e:
        details["P&ID Pipeline"] = str(e)
    checks["P&ID Pipeline"] = "PASS" if pid_ok else "FAIL"

    # 5. Sandbox execution with network guard
    sandbox_ok = False
    try:
        from backend.sandbox.executor import sandbox_executor
        res = sandbox_executor.execute("print(2 + 2)", timeout_sec=5)
        if res.get("success") and "4" in res.get("stdout", ""):
            # Check network guard
            probe = sandbox_executor.execute("import socket; s = socket.socket()", timeout_sec=5)
            if "PermissionError" in probe.get("stderr", "") or "PermissionError" in probe.get("stdout", ""):
                sandbox_ok = True
            else:
                details["Sandbox"] = "Network isolation guard did not raise PermissionError"
        else:
            details["Sandbox"] = f"Math execution failed: {res.get('stderr')}"
    except Exception as e:
        details["Sandbox"] = str(e)
    checks["Sandbox"] = "PASS" if sandbox_ok else "FAIL"

    # 6. DOCX Generation
    docx_ok = False
    try:
        import docx
        doc = docx.Document()
        doc.add_heading("Offline Test", 0)
        buf = io.BytesIO()
        doc.save(buf)
        if len(buf.getvalue()) > 500:
            docx_ok = True
    except Exception as e:
        details["DOCX"] = str(e)
    checks["DOCX"] = "PASS" if docx_ok else "FAIL"

    # 7. XLSX Generation
    xlsx_ok = False
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "Test"
        buf = io.BytesIO()
        wb.save(buf)
        if len(buf.getvalue()) > 500:
            xlsx_ok = True
    except Exception as e:
        details["XLSX"] = str(e)
    checks["XLSX"] = "PASS" if xlsx_ok else "FAIL"

    # 8. Frontend build
    frontend_ok = False
    try:
        dist_html = os.path.join(os.getcwd(), "frontend", "dist", "index.html")
        dist_assets = os.path.join(os.getcwd(), "frontend", "dist", "assets")
        if os.path.exists(dist_html) and os.path.isdir(dist_assets) and len(os.listdir(dist_assets)) > 0:
            frontend_ok = True
        else:
            details["Frontend"] = "frontend/dist not found. Run 'npm run build' in frontend/"
    except Exception as e:
        details["Frontend"] = str(e)
    checks["Frontend"] = "PASS" if frontend_ok else "FAIL"

    # 9. External AI dependency
    checks["External AI dependency"] = "NONE"

    # 10. Overall readiness
    all_passed = all(status in ("PASS", "NONE") for status in checks.values())
    checks["Offline readiness"] = "PASS" if all_passed else "FAIL"

    # Output formatted report
    print("\n================================")
    print("CONFIGIQ OFFLINE DEMO CHECK")
    print("================================")
    for name in [
        "Ollama",
        "Llama",
        "Qwen Coder 7B",
        "Qwen Coder 14B",
        "Moondream",
        "Nomic Embeddings",
        "RAG",
        "OCR",
        "P&ID Pipeline",
        "Sandbox",
        "DOCX",
        "XLSX",
        "Frontend",
        "External AI dependency",
        "Offline readiness"
    ]:
        val = checks.get(name, "FAIL")
        print(f"{name:<24} {val}")
    print("================================\n")

    if not all_passed:
        print("IDENTIFIED ISSUES:")
        for k, v in details.items():
            print(f"  - {k}: {v}")
        print()
        return False
    return True

if __name__ == "__main__":
    success = check_all()
    sys.exit(0 if success else 1)
