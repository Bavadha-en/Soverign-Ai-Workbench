import os
import sys
import io
import pytest
from PIL import Image
from fastapi.testclient import TestClient
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.main import app
from backend.agents.agent import agent_orchestrator
from backend.agents.schemas import AgentStatus
from backend.rag.ingest import ingestion_engine
from backend.rag.retriever import retriever
from backend.services.network_monitor import network_monitor
from backend.sandbox.executor import sandbox_executor
from backend.llm.model_router import model_router

client = TestClient(app)

@pytest.fixture(scope="module", autouse=True)
def init_kb():
    kb_dir = os.path.join(os.getcwd(), "knowledge_base")
    if os.path.isdir(kb_dir):
        ingestion_engine.ingest_directory(kb_dir, force_reindex=False)


# 1. Inspection -> Approval Note
@pytest.mark.asyncio
async def test_workflow_a_inspection_to_approval_note():
    doc_path = os.path.abspath("demo_data/inspection/inspection_report_P101_clean.pdf")
    task = "Inspect centrifugal pump P-101 for vibration and bearing temperature, check SOP-M-104, and generate approval note"
    state = await agent_orchestrator.run(
        task=task,
        document_ids=[doc_path],
        parameters={"file_path": doc_path}
    )

    assert state.status == AgentStatus.COMPLETED
    assert len(state.completed_steps) >= 5
    assert len(state.generated_files) >= 1

    docx_file = state.generated_files[0]
    assert docx_file.endswith(".docx")
    assert os.path.exists(docx_file)

    doc = Document(docx_file)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any("Reference Document" in h for h in headings)
    assert any("Executive Summary" in h for h in headings)
    assert any("Inspection Findings" in h for h in headings)
    assert any("SOP" in h for h in headings)
    assert any("Risk" in h for h in headings)
    assert any("Recommended Actions" in h for h in headings)
    assert any("Approval Recommendation" in h for h in headings)
    assert any("Sources" in h for h in headings)


# 2. RAG source grounding
def test_workflow_e_rag_source_grounding():
    results = retriever.retrieve("centrifugal pump elevated vibration ISO 10816 SOP", top_k=3)
    assert len(results) >= 1
    for r in results:
        assert "document" in r
        assert r["document"] != ""
        assert "page" in r
        assert "score" in r
        assert isinstance(r["score"], float)
        assert "content" in r
        assert len(r["content"]) > 10


# 3. P&ID -> VLM -> RAG -> answer
@pytest.mark.asyncio
async def test_workflow_b_pid_vlm_rag_answer():
    # Self-made diagram (scripts/generate_pid_b.py), safe to ship
    pid_path = os.path.abspath("demo_data/pid/pid_system_b.png")
    assert os.path.exists(pid_path), f"P&ID image not found at {pid_path}"

    task = "Analyze P&ID diagram for safety relief valve and pressure rating per pressure vessel SOP"
    state = await agent_orchestrator.run(
        task=task,
        document_ids=["doc_pid_01"],
        parameters={"file_path": pid_path}
    )

    assert state.status == AgentStatus.COMPLETED
    assert "extract_document" in state.tool_results
    assert "analyze_scanned_pages" in state.tool_results
    assert "search_maintenance_sop" in state.tool_results
    assert "VISUAL EVIDENCE" in state.final_output
    assert "DOCUMENT EVIDENCE" in state.final_output


# 4. Coding -> sandbox -> verification
@pytest.mark.asyncio
async def test_workflow_c_coding_sandbox_verification():
    calc_task = "Calculate pump hydraulic efficiency for flow rate 50 m3/h, head 60 m, power 11 kW"
    state = await agent_orchestrator.run(task=calc_task)

    assert state.status == AgentStatus.COMPLETED
    assert "execute_in_sandbox" in state.tool_results
    sandbox_res = state.tool_results["execute_in_sandbox"]
    assert sandbox_res.get("exit_code") == 0
    assert "stdout" in sandbox_res
    assert len(sandbox_res["stdout"]) > 0
    assert "Efficiency" in sandbox_res["stdout"] or "Hydraulic Power" in sandbox_res["stdout"]


# 5. Network telemetry
def test_workflow_d_network_telemetry_sovereignty():
    telemetry = network_monitor.get_telemetry()
    assert telemetry.status in ("LOCAL_ONLY", "WARNING_EXTERNAL_ATTEMPT_DETECTED")
    assert len(telemetry.active_listening_ports) >= 1
    assert "SOVEREIGN-SEAL-" in telemetry.integrity_hash

    # Test sandbox blocked network execution
    blocked_script = "import socket\nsocket.socket(socket.AF_INET, socket.SOCK_STREAM)"
    res = sandbox_executor.execute(blocked_script)
    assert res["exit_code"] != 0 or "PermissionError" in res["stderr"] or "forbidden" in res["stderr"].lower()


# 6. Model routing
def test_workflow_f_model_routing():
    res_code = model_router.route("Write Python code to compute pipe Reynolds number")
    assert "qwen" in res_code["model"].lower() or "coder" in res_code["model"].lower()

    res_vision = model_router.route("Inspect this visual scanned image", image_present=True)
    assert "moondream" in res_vision["model"].lower()

    res_gen = model_router.route("Explain safety protocols for boiler blowdown operation")
    assert "llama" in res_gen["model"].lower()


# 7. DOCX dynamic-content validation (Anti-Hardcode Check)
@pytest.mark.asyncio
async def test_workflow_g_docx_dynamic_anti_hardcode():
    doc_a = os.path.abspath("demo_data/inspection/inspection_report_P101_clean.pdf")
    doc_b = os.path.abspath("demo_data/inspection/inspection_report_P202_clean.pdf")

    state_a = await agent_orchestrator.run(
        task="Inspect centrifugal pump P-101",
        document_ids=[doc_a],
        parameters={"file_path": doc_a}
    )
    state_b = await agent_orchestrator.run(
        task="Inspect secondary booster pump P-202",
        document_ids=[doc_b],
        parameters={"file_path": doc_b}
    )

    docx_a = state_a.generated_files[0]
    docx_b = state_b.generated_files[0]

    doc_a_obj = Document(docx_a)
    doc_b_obj = Document(docx_b)

    text_a = " ".join([p.text for p in doc_a_obj.paragraphs])
    text_b = " ".join([p.text for p in doc_b_obj.paragraphs])

    assert "P-101" in text_a
    assert "P-202" in text_b
    assert "P-202" not in text_a
    assert "P-101" not in text_b
    assert text_a != text_b
