"""
Accuracy of the inspection report -> approval note flow: values are read
exactly, checked against the SOP the report cites, and anything doubtful is
flagged for the engineer instead of passed through.
"""
import os
import sys

import pytest
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.agents.agent import ConfigIQAgent
from backend.agents.schemas import FactVerificationStatus
from backend.agents.verifier import verifier
from backend.documents.inspection_checks import assess_document, assess_text
from backend.documents.ocr import ocr_engine

INSPECTION = os.path.join("demo_data", "inspection")
needs_ocr = pytest.mark.skipif(not ocr_engine.is_available(), reason="RapidOCR is not installed")


def _report_text(name: str = "inspection_report_P101_clean.txt") -> str:
    with open(os.path.join(INSPECTION, name), encoding="utf-8") as f:
        return f.read()


def _by_label(assessment):
    return {m["label_full"]: m for m in assessment["measurements"]}


def test_narrative_report_values_and_checks():
    a = assess_text(_report_text())
    assert a["equipment_tag"]["value"] == "P-101"
    assert a["inspection_date"]["value"] == "2026-08-14"

    m = _by_label(a)
    outboard = m["Vibration (Outboard Bearing RMS)"]
    assert (outboard["value"], outboard["unit"], outboard["status"], outboard["zone"]) == (7.4, "mm/s", "CRITICAL", "D")
    assert outboard["limit_source"] == "SOP-M-104 §3"
    assert m["Vibration (Drive End Bearing RMS)"]["status"] == "WARNING"
    assert (m["Bearing Temperature"]["value"], m["Bearing Temperature"]["status"]) == (88.0, "EXCEEDED")
    assert (m["Mechanical Seal Dynamic Leakage"]["value"], m["Mechanical Seal Dynamic Leakage"]["status"]) == (14.0, "EXCEEDED")
    assert m["Flow Rate"]["status"] == "DEVIATION"

    assert a["severity"]["value"] == "CRITICAL"
    assert a["severity"]["agrees"] is True
    assert a["review_items"] == []


def test_table_layout_pdf_is_read():
    a = assess_document(os.path.join(INSPECTION, "inspection_report_P101_clean.pdf"))
    assert a["text_source"] == "text_layer"
    m = {x["kind"]: x for x in a["measurements"]}
    assert (m["vibration"]["value"], m["vibration"]["status"]) == (7.4, "CRITICAL")
    assert (m["suction_pressure"]["basis"], m["suction_pressure"]["status"]) == ("report_range", "OK")
    assert (m["flow_rate"]["basis"], m["flow_rate"]["status"]) == ("report_reference", "DEVIATION")
    assert a["severity"]["value"] == "CRITICAL"
    # This form of the report has no date, and says so instead of inventing one.
    assert a["review_items"] == ["The report does not state an inspection date."]


def test_severity_disagreement_is_flagged():
    a = assess_text(_report_text().replace("Overall Risk Severity: CRITICAL", "Overall Risk Severity: LOW"))
    assert a["severity"]["value"] == "CRITICAL"
    assert any("rates severity LOW" in item for item in a["review_items"])


def test_missing_required_reading_is_flagged():
    text = "\n".join(line for line in _report_text().splitlines() if "Seal Dynamic Leakage" not in line)
    a = assess_text(text)
    assert not any(m["kind"] == "seal_leakage" for m in a["measurements"])
    assert any("seal leakage reading" in item for item in a["review_items"])


def test_findings_value_that_disagrees_with_the_table_is_flagged():
    a = assess_text(_report_text().replace("bearing cap temperature of 88.0 deg C", "bearing cap temperature of 8.0 deg C"))
    assert any("8.0 °C" in item and "88" in item for item in a["review_items"])


def test_report_quoting_a_different_limit_is_flagged():
    a = assess_text(_report_text().replace("max allowable is 75.0 deg C", "max allowable is 90.0 deg C"))
    assert any("bearing temperature limit of 90" in item for item in a["review_items"])


def test_verifier_rejects_a_figure_that_is_not_in_the_report():
    summary = verifier.verify_facts(
        [
            "Bearing temperature reached 88.0 deg C on the outboard bearing.",
            "Bearing temperature reached 68.0 deg C on the outboard bearing.",
        ],
        [],
        document_text=_report_text(),
    )
    right, wrong = summary.claims
    assert right.status != FactVerificationStatus.UNSUPPORTED
    assert wrong.status == FactVerificationStatus.UNSUPPORTED
    assert "68 °C" in wrong.evidence


@needs_ocr
def test_scanned_image_reads_every_required_value():
    a = assess_document(os.path.join(INSPECTION, "inspection_report_P101_scanned.png"))
    assert a["text_source"] == "ocr"
    values = {m["kind"]: m["value"] for m in a["measurements"] if m["kind"] != "vibration"}
    assert values["bearing_temperature"] == 88.0
    assert values["seal_leakage"] == 14.0
    assert max(m["value"] for m in a["measurements"] if m["kind"] == "vibration") == 7.4
    assert a["severity"]["value"] == "CRITICAL"


@needs_ocr
def test_scanned_pdf_is_read_with_ocr():
    a = assess_document(os.path.join(INSPECTION, "inspection_report_P202_scanned.pdf"))
    assert a["text_source"] == "ocr"
    assert a["equipment_tag"]["value"] == "P-202"
    assert a["inspection_date"]["value"] == "2026-08-22"
    assert a["severity"]["value"] == "HIGH"


@pytest.mark.asyncio
async def test_approval_note_is_built_from_checked_values():
    path = os.path.abspath(os.path.join(INSPECTION, "inspection_report_P101_clean.pdf"))
    state = await ConfigIQAgent().run(
        task="Review this inspection report and prepare the approval note.",
        document_ids=[path],
        parameters={"file_path": path},
    )
    assert state.tool_results["check_readings"]["applicable"]

    doc = Document(next(p for p in state.generated_files if p.endswith(".docx")))
    paragraphs = [p.text for p in doc.paragraphs]
    rows = [" | ".join(c.text for c in row.cells) for t in doc.tables for row in t.rows]

    assert "Risk Level: CRITICAL" in paragraphs
    outboard = next(r for r in rows if r.startswith("Vibration RMS (Outboard)"))
    assert "7.4 mm/s" in outboard and "Critical (Zone D)" in outboard and "SOP-M-104 §3" in outboard
    assert "Review item: The report does not state an inspection date." in paragraphs
    # A note with open review items never claims to be verified.
    assert not any("SOURCED & VERIFIED" in p for p in paragraphs)
