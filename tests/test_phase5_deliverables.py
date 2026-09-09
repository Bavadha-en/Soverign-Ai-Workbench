import os
import sys
import pytest
from fastapi.testclient import TestClient
from docx import Document
import openpyxl
from pptx import Presentation

# Ensure ConfigIQ root is on PYTHONPATH
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.main import app
from backend.agents.schemas import AgentStatus
from backend.agents.tool_registry import ToolRegistry, tool_registry
from backend.agents.agent import ConfigIQAgent
from backend.agents.planner import Planner
from backend.tools.word_tool import create_approval_note_docx
from backend.tools.excel_tool import create_calculation_xlsx
from backend.tools.ppt_tool import create_executive_summary_pptx

client = TestClient(app)


# 1. DOCX Generation & File Creation
def test_docx_generation():
    out_path = os.path.join("outputs", "test_deliverable_01.docx")
    if os.path.exists(out_path):
        os.remove(out_path)

    res_path = create_approval_note_docx(
        output_path=out_path,
        task_id="task_docx_test_01",
        reference_document="Control_Valve_CV102_Report.pdf",
        executive_summary="Critical wall loss identified; replacement required under SOP-M-402."
    )
    assert os.path.exists(res_path)
    assert res_path.endswith(".docx")
    assert os.path.getsize(res_path) > 1000


# 2. XLSX Generation & File Creation
def test_xlsx_generation():
    out_path = os.path.join("outputs", "test_deliverable_02.xlsx")
    if os.path.exists(out_path):
        os.remove(out_path)

    res_path = create_calculation_xlsx(
        output_path=out_path,
        task_id="task_xlsx_test_02",
        title="Pump Efficiency Calculation"
    )
    assert os.path.exists(res_path)
    assert res_path.endswith(".xlsx")
    assert os.path.getsize(res_path) > 1000


# 3. PPTX Generation & File Creation
def test_pptx_generation():
    out_path = os.path.join("outputs", "test_deliverable_03.pptx")
    if os.path.exists(out_path):
        os.remove(out_path)

    res_path = create_executive_summary_pptx(
        output_path=out_path,
        task_id="task_pptx_test_03",
        reference_document="Control_Valve_CV102_Report.pdf"
    )
    assert os.path.exists(res_path)
    assert res_path.endswith(".pptx")
    assert os.path.getsize(res_path) > 5000


# 4. DOCX Required 8 Sections Verification
def test_docx_required_sections():
    out_path = os.path.join("outputs", "test_docx_sections.docx")
    create_approval_note_docx(
        output_path=out_path,
        task_id="task_docx_sec_04",
        reference_document="Inspection_Report_CV102.pdf",
        executive_summary="Executive summary content for test.",
        inspection_findings=["Wall loss 36% measured via UT."],
        sop_references=["SOP-M-402: High Pressure Control Valve Replacement"],
        risk_severity="HIGH",
        recommended_actions=["Mandatory replacement with 316L valve."],
        approval_recommendation="APPROVED FOR WORK ORDER",
        sources=[{"document": "SOP-M-402.pdf", "page": 3, "score": 0.95, "status": "SUPPORTED"}],
        model_used="llama3:latest",
        verification_status="SUPPORTED"
    )
    doc = Document(out_path)
    full_text = " ".join([p.text for p in doc.paragraphs if p.text])

    # 8 Required Sections
    assert "1. Reference Document" in full_text
    assert "2. Executive Summary" in full_text
    assert "3. Inspection Findings" in full_text
    assert "4. SOP / Manual References" in full_text
    assert "5. Risk / Severity Assessment" in full_text
    assert "6. Recommended Actions" in full_text
    assert "7. Approval Recommendation" in full_text
    assert "8. Verified Sources & Knowledge Provenance" in full_text

    # Metadata & Source references
    assert "Task ID: task_docx_sec_04" in full_text
    assert "llama3:latest" in full_text
    assert "STATUS: SUPPORTED" in full_text.upper()
    assert "Page 3" in full_text


# 5. XLSX Required 4 Sheets Verification
def test_xlsx_required_sheets():
    out_path = os.path.join("outputs", "test_xlsx_sheets.xlsx")
    create_calculation_xlsx(
        output_path=out_path,
        task_id="task_xlsx_sheets_05",
        title="Pump Efficiency Verification"
    )
    wb = openpyxl.load_workbook(out_path)
    sheet_names = wb.sheetnames
    assert "Inputs" in sheet_names
    assert "Calculation" in sheet_names
    assert "Verification" in sheet_names
    assert "Sources" in sheet_names

    # Check Inputs Sheet Columns
    ws_inputs = wb["Inputs"]
    headers_inputs = [cell.value for cell in ws_inputs[4] if cell.value]
    assert "Parameter" in headers_inputs
    assert "Value" in headers_inputs
    assert "Unit" in headers_inputs
    assert "Source" in headers_inputs

    # Check Calculation Sheet Columns
    ws_calc = wb["Calculation"]
    headers_calc = [cell.value for cell in ws_calc[4] if cell.value]
    assert "Step / Parameter" in headers_calc
    assert "Formula" in headers_calc
    assert "Substitution" in headers_calc
    assert "Final Result" in headers_calc

    # Check Verification Sheet Columns
    ws_verif = wb["Verification"]
    headers_verif = [cell.value for cell in ws_verif[4] if cell.value]
    assert "Check" in headers_verif
    assert "Result" in headers_verif
    assert "Status" in headers_verif

    # Check Sources Sheet Columns
    ws_src = wb["Sources"]
    headers_src = [cell.value for cell in ws_src[4] if cell.value]
    assert "Input / Finding" in headers_src
    assert "Source Type" in headers_src
    assert "Verification Status" in headers_src


# 6. PPTX 8 Slides Structure & Content
def test_pptx_slide_count_and_structure():
    out_path = os.path.join("outputs", "test_pptx_slides.pptx")
    create_executive_summary_pptx(
        output_path=out_path,
        task_id="task_pptx_06",
        reference_document="Inspection_Report_CV102.pdf",
        title="Inspection Report Review"
    )
    prs = Presentation(out_path)
    assert len(prs.slides) == 8

    # Check aspect ratio 16:9 widescreen
    assert round(prs.slide_width.inches, 2) == 13.33
    assert round(prs.slide_height.inches, 2) == 7.5


# 7. Source Metadata Traceability
def test_source_metadata_traceability():
    docx_path = os.path.join("outputs", "test_source_meta.docx")
    sources = [
        {"document": "Inspection_Report_CV102.pdf", "page": 4, "score": 0.98, "status": "SUPPORTED"},
        {"document": "SOP-M-402.pdf", "page": 2, "score": 0.92, "status": "SUPPORTED"}
    ]
    create_approval_note_docx(
        output_path=docx_path,
        task_id="task_source_07",
        reference_document="Inspection_Report_CV102.pdf",
        sources=sources
    )
    doc = Document(docx_path)
    text = " ".join([p.text for p in doc.paragraphs])
    assert "Page 4" in text
    assert "Page 2" in text
    assert "Inspection_Report_CV102.pdf" in text
    assert "SOP-M-402.pdf" in text


# 8. Verification Status Embedding
def test_verification_status_embedding():
    docx_path = os.path.join("outputs", "test_verif_status.docx")
    create_approval_note_docx(
        output_path=docx_path,
        task_id="task_verif_08",
        reference_document="Report.pdf",
        verification_status="SUPPORTED"
    )
    doc = Document(docx_path)
    text = " ".join([p.text for p in doc.paragraphs])
    assert "SUPPORTED" in text.upper()
    assert "SOURCED & VERIFIED" in text.upper()


# 9. Human Review Flag in Word when Unsupported / Needs Review
def test_human_review_flag_triggering():
    docx_path = os.path.join("outputs", "test_human_review.docx")
    create_approval_note_docx(
        output_path=docx_path,
        task_id="task_review_09",
        reference_document="Unverified_Report.pdf",
        verification_status="NEEDS REVIEW",
        human_review_required=True
    )
    doc = Document(docx_path)
    text = " ".join([p.text for p in doc.paragraphs])
    assert "HUMAN REVIEW REQUIRED" in text.upper()
    assert "AI-assisted draft" in text


# 10. Multiple Generated Files in Single Task
@pytest.mark.asyncio
async def test_multiple_generated_files():
    agent = ConfigIQAgent()
    state = await agent.run(
        task="Analyze the inspection report and prepare all management deliverables.",
        document_ids=["doc_multi_test"]
    )
    assert state.status == AgentStatus.COMPLETED
    assert len(state.generated_files) >= 2

    extensions = [os.path.splitext(f)[1].lower() for f in state.generated_files]
    assert ".docx" in extensions
    assert ".pptx" in extensions

    for fpath in state.generated_files:
        assert os.path.exists(fpath)


# 11. ToolRegistry Registration & Listing
def test_tool_registry_phase5_tools():
    reg = ToolRegistry()
    tool_names = [t.name for t in reg.list_tools()]

    # Verify all expected tools are present
    assert "document_reader" in tool_names
    assert "pdf_processor" in tool_names
    assert "ocr" in tool_names
    assert "vision" in tool_names
    assert "rag_search" in tool_names
    assert "llm_generate" in tool_names
    assert "code_executor" in tool_names
    assert "document_generator" in tool_names
    assert "verification" in tool_names
    assert "excel_generator" in tool_names
    assert "ppt_generator" in tool_names

    resp = client.get("/agent/tools")
    assert resp.status_code == 200
    api_tools = [t["name"] for t in resp.json()]
    assert "excel_generator" in api_tools
    assert "ppt_generator" in api_tools
    assert "document_generator" in api_tools


# 12. Agent Integration with Dynamic Tool Selection
def test_agent_dynamic_planner_selection():
    p = Planner()

    # Calculation request -> Excel
    calc_plan = p.plan("Calculate pump efficiency and create calculation workbook.")
    calc_tools = [s["tool"] for s in calc_plan]
    assert "excel_generator" in calc_tools

    # Executive presentation request -> PPTX
    ppt_plan = p.plan("Prepare an executive presentation for the inspection report.")
    ppt_tools = [s["tool"] for s in ppt_plan]
    assert "ppt_generator" in ppt_tools

    # All deliverables request -> DOCX + PPTX + XLSX
    all_plan = p.plan("Analyze the inspection report and prepare all management deliverables.")
    all_tools = [s["tool"] for s in all_plan]
    assert "document_generator" in all_tools
    assert "ppt_generator" in all_tools


# 13. Output Path Safety
def test_output_path_safety():
    out_dir = os.path.abspath("outputs")
    docx_path = create_approval_note_docx(
        output_path=os.path.join("outputs", "test_path_safety.docx"),
        task_id="task_safe_13"
    )
    abs_gen = os.path.abspath(docx_path)
    assert abs_gen.startswith(out_dir)


# 14. Path Traversal Protection on Output Download Endpoint
def test_path_traversal_protection():
    # 1. Reject parent directory traversal
    resp_trav1 = client.get("/outputs/../main.py")
    assert resp_trav1.status_code in [400, 403, 404]

    # 2. Reject encoded traversal
    resp_trav2 = client.get("/outputs/..%2F..%2Fetc%2Fpasswd")
    assert resp_trav2.status_code in [400, 403, 404]

    # 3. Reject absolute Windows paths
    resp_trav3 = client.get("/outputs/C:/Windows/System32/drivers/etc/hosts")
    assert resp_trav3.status_code in [400, 403, 404]

    # 4. Valid file download succeeds
    valid_file = os.path.join("outputs", "test_download_valid.txt")
    with open(valid_file, "w") as f:
        f.write("Valid output file content")
    resp_valid = client.get("/outputs/test_download_valid.txt")
    assert resp_valid.status_code == 200
    assert resp_valid.text == "Valid output file content"


# 15. No External Network Dependency Verification
def test_no_external_network_dependency():
    health_resp = client.get("/health")
    assert health_resp.status_code == 200
    data = health_resp.json()
    assert data["network"] == "LOCAL_ONLY"
    assert data["environment"] == "on-premise / air-gapped"


# 16. End-to-End Demo 1: Inspection Report -> DOCX + PPTX
@pytest.mark.asyncio
async def test_end_to_end_demo1_inspection_to_docx_and_pptx():
    agent = ConfigIQAgent()
    state = await agent.run(
        task="Analyze this inspection report and prepare an approval note and executive summary.",
        document_ids=["doc_valve_inspection_report"]
    )
    assert state.status == AgentStatus.COMPLETED
    assert len(state.completed_steps) >= 6
    assert len(state.generated_files) >= 2

    # Check generated files
    has_docx = any(f.endswith(".docx") for f in state.generated_files)
    has_pptx = any(f.endswith(".pptx") for f in state.generated_files)
    assert has_docx, "Expected Word (.docx) approval note to be generated"
    assert has_pptx, "Expected PowerPoint (.pptx) executive summary to be generated"

    for f in state.generated_files:
        assert os.path.exists(f)
        assert os.path.getsize(f) > 0


# 17. End-to-End Demo 2: Engineering Calculation in Sandbox -> XLSX Workbook
@pytest.mark.asyncio
async def test_end_to_end_demo2_calculation_to_xlsx():
    agent = ConfigIQAgent()
    state = await agent.run(
        task="Calculate pump efficiency using flow rate of 50 m3/h, head of 60 m, and power of 11 kW and create a calculation workbook."
    )
    assert state.status == AgentStatus.COMPLETED
    assert "execute_in_sandbox" in state.tool_results

    # Verify calculation output in stdout
    sandbox_res = state.tool_results["execute_in_sandbox"]
    stdout = sandbox_res["stdout"]
    assert "Hydraulic Power" in stdout
    assert "Efficiency" in stdout

    # Verify XLSX workbook was generated
    has_xlsx = any(f.endswith(".xlsx") for f in state.generated_files)
    assert has_xlsx, "Expected Excel (.xlsx) calculation workbook to be generated"

    xlsx_path = [f for f in state.generated_files if f.endswith(".xlsx")][0]
    assert os.path.exists(xlsx_path)

    # Validate Workbook contents
    wb = openpyxl.load_workbook(xlsx_path)
    assert "Calculation" in wb.sheetnames
    ws_calc = wb["Calculation"]
    calc_values = [cell.value for row in ws_calc.iter_rows() for cell in row if cell.value is not None]

    # Verify Hydraulic Power ~ 8.175 kW and Efficiency ~ 74.32%
    assert any("8.175" in str(v) or v == 8.175 for v in calc_values)
    assert any("74.32" in str(v) or v == 74.32 for v in calc_values)
