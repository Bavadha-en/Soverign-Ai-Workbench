import os
import sys
import pytest
from docx import Document

# Ensure ConfigIQ root is on PYTHONPATH
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from backend.tools.word_tool import generate_approval_note, create_word
from backend.services.audit_service import audit_service


def test_phase8_approval_note_generation(tmp_path):
    output_docx = str(tmp_path / "Approval_Note.docx")

    res = generate_approval_note(
        output_path=output_docx,
        title="APPROVAL NOTE - INDUSTRIAL EQUIPMENT REPAIR",
        subject="Request for Approval to Replace Corroded Valve Assembly CV-102",
        background="Inspection conducted on 2026-09-08 identified severe surface corrosion on CV-102.",
        inspection_findings="Ultrasonic thickness gauging indicates wall thickness reduced to 4.2mm (threshold 8.0mm).",
        technical_assessment="Valve operates at 25 BAR. Risk of catastrophic wall rupture if not replaced immediately.",
        applicable_sop="SOP-MECH-44: High-Pressure Piping & Valve Replacement Guidelines.",
        recommended_action="Isolate Line-4B during planned shutdown and replace CV-102 with standard alloy assembly.",
        approval_requested="Approval requested for emergency procurement of CV-102 replacement assembly.",
        task_id="task_phase8_test"
    )

    assert res["success"] is True
    assert res["output_path"] == output_docx
    assert res["file_size"] > 0
    assert os.path.exists(output_docx)

    # Verify actual .docx document structure and content
    doc = Document(output_docx)
    all_text = "\n".join([p.text for p in doc.paragraphs])

    assert "APPROVAL NOTE" in all_text
    assert "Subject:" in all_text
    assert "1. Background" in all_text
    assert "2. Inspection Findings" in all_text
    assert "3. Technical Assessment" in all_text
    assert "4. Applicable SOP" in all_text
    assert "5. Recommended Action" in all_text
    assert "6. Approval Requested" in all_text
    assert "SOP-MECH-44" in all_text

    # Verify sign-off table
    assert len(doc.tables) >= 1
    table_text = "\n".join([cell.text for row in doc.tables[0].rows for cell in row.cells])
    assert "Role" in table_text
    assert "Prepared By" in table_text
    assert "Approved By" in table_text

    # Verify audit log
    logs = audit_service.get_logs(task_id="task_phase8_test")
    assert len(logs) > 0
    assert logs[0].action == "GENERATE_APPROVAL_NOTE"
    assert logs[0].status == "SUCCESS"


def test_phase8_default_output():
    default_path = "outputs/Approval_Note.docx"
    if os.path.exists(default_path):
        os.remove(default_path)

    res = generate_approval_note()
    assert res["success"] is True
    assert os.path.exists(default_path)
    assert res["file_size"] > 0
